# -*- coding: utf-8 -*-
"""Generate fillable DOCX: Feedback & Self-Tracking Form for TM-002 after Session 2."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

OUT = Path(__file__).parent / "Feedback-Form-After-S2-TM-002.docx"

doc = Document()

for section in doc.sections:
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

style = doc.styles["Normal"]
style.font.name = "Segoe UI"
style.font.size = Pt(10.5)
style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
style.paragraph_format.space_after = Pt(4)
style.paragraph_format.line_spacing = 1.15


def heading1(text):
    p = doc.add_heading(text, level=1)
    for run in p.runs:
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
    return p


def heading2(text):
    p = doc.add_heading(text, level=2)
    for run in p.runs:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
    p.paragraph_format.space_before = Pt(14)
    return p


def question(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    return p


def instruction(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    run.italic = True
    return p


def body_text(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    return p


def fill_line(count=1):
    for _ in range(count):
        p = doc.add_paragraph("_" * 85)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(6)
        for run in p.runs:
            run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
            run.font.size = Pt(10)


def fill_box(lines=3):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("\n" * lines)
    run.font.size = Pt(10)
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.makeelement(qn("w:pBdr"), {})
    for side in ["top", "left", "bottom", "right"]:
        bdr = pBdr.makeelement(
            qn("w:{}".format(side)),
            {
                qn("w:val"): "single",
                qn("w:sz"): "4",
                qn("w:space"): "4",
                qn("w:color"): "CCCCCC",
            },
        )
        pBdr.append(bdr)
    pPr.append(pBdr)


def checkbox_line(text, bold_prefix=None):
    p = doc.add_paragraph()
    run = p.add_run("\u2610  ")
    run.font.size = Pt(12)
    if bold_prefix:
        rb = p.add_run(bold_prefix)
        rb.bold = True
        rb.font.size = Pt(10.5)
    run2 = p.add_run(text)
    run2.font.size = Pt(10.5)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    return p


def set_cell_shading(cell, color):
    shading = cell._tc.get_or_add_tcPr().makeelement(
        qn("w:shd"),
        {qn("w:fill"): color, qn("w:val"): "clear"},
    )
    cell._tc.get_or_add_tcPr().append(shading)


def styled_table(headers, rows, col_widths=None):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        set_cell_shading(cell, "F0F4F8")
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.rows[ri + 1].cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.size = Pt(10)
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in tbl.rows:
                row.cells[i].width = Cm(w)
    return tbl


# ============================================================
# DOCUMENT
# ============================================================

heading1("Feedback & Self-Tracking Form")
meta = doc.add_paragraph()
run = meta.add_run(
    "T.M. (TM-002)  |  After Session 2 (Oct 16, 2026)  |  "
    "Sessions completed: 2 of 6  |  "
    "Complete before Session 3  |  "
    "Be honest \u2014 no right or wrong answers"
)
run.font.size = Pt(9)
run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

instruction(
    "This form has two purposes: (1) give me feedback on how our session landed, "
    "and (2) help you track your own progress day by day between now and our next session. "
    "Type directly into this document or print and fill in by hand."
)

# ===== SECTION 1 =====
heading2("Section 1: Overall Check-In")
question("Since our session on Oct 16, overall I feel:")
for label in [
    "Much Worse",
    "Somewhat Worse",
    "About the Same",
    "Somewhat Better",
    "Much Better",
]:
    checkbox_line(label)

question("In your own words, what has changed (if anything)?")
fill_box(2)

# ===== SECTION 2 =====
heading2("Section 2: Where I Am Right Now (1\u201310 Scales)")
instruction(
    "Rate each area honestly. Same scales we will use at the start of every session. "
    "There are no wrong answers \u2014 this is your personal baseline."
)
styled_table(
    ["Area", "What the scale means", "Your Rating"],
    [
        ["Self-Esteem", "1 = no value in who I am \u2026 10 = confident and worthy", "      /10"],
        ["Assertiveness", "1 = cannot speak up \u2026 10 = comfortably express needs", "      /10"],
        ["Anxiety (calm)", "1 = anxiety runs the show \u2026 10 = calm and in control", "      /10"],
        ["Trust", "1 = guard always up \u2026 10 = open with chosen people", "      /10"],
        ["Comfort in Own Skin", "1 = deeply uncomfortable \u2026 10 = at ease and at home", "      /10"],
        ["Influence of the Past", "1 = past controls everything \u2026 10 = aware but free", "      /10"],
    ],
    col_widths=[3.5, 9, 2.5],
)
question("Which area feels most different compared to before our sessions started?")
fill_line(1)

# ===== SECTION 3 =====
heading2("Section 3: Daily Self-Tracking Log")
instruction(
    "Fill in one row each day (or as often as you can). "
    "This is for YOUR awareness \u2014 noticing patterns is half the work."
)
styled_table(
    [
        "Day / Date",
        "Tribe Viz\n(before sleep)",
        "Anchor\n(thumb+finger)",
        "Breathing\nAffirmation",
        "Rescuer-Helper\nReframe",
        "Session\nRecording",
        "Mood\n(1\u201310)",
        "Notes / What happened",
    ],
    [
        ["Day 1:", "\u2610", "\u2610", "\u2610", "\u2610", "\u2610", "  /10", ""],
        ["Day 2:", "\u2610", "\u2610", "\u2610", "\u2610", "\u2610", "  /10", ""],
        ["Day 3:", "\u2610", "\u2610", "\u2610", "\u2610", "\u2610", "  /10", ""],
        ["Day 4:", "\u2610", "\u2610", "\u2610", "\u2610", "\u2610", "  /10", ""],
        ["Day 5:", "\u2610", "\u2610", "\u2610", "\u2610", "\u2610", "  /10", ""],
        ["Day 6:", "\u2610", "\u2610", "\u2610", "\u2610", "\u2610", "  /10", ""],
        ["Day 7:", "\u2610", "\u2610", "\u2610", "\u2610", "\u2610", "  /10", ""],
    ],
    col_widths=[2, 1.8, 1.8, 1.8, 1.8, 1.8, 1.3, 3.7],
)
instruction(
    "Practices: Tribe Viz = picture the 11-member circle before sleep. "
    "Anchor = press thumb + forefinger (right hand) when you need steadiness. "
    "Breathing = \"I am [word 1] and I am [word 2] and I like it\" (3 breaths, morning + bedtime). "
    "Rescuer-Helper = reframe someone's anger as their struggle. "
    "Recording = listened to the session audio."
)

# ===== SECTION 4 =====
heading2("Section 4: The Anchor")
question(
    "When you press your thumb and forefinger together (right hand), what happens?"
)
checkbox_line("I feel the steadiness / protection / \"armor\" return")
checkbox_line("I feel something but it is faint")
checkbox_line("I do not feel much yet")
checkbox_line("I have not tried it")
question(
    "Describe one specific moment you used the anchor this week. "
    "Where were you? What was happening? What did it feel like?"
)
fill_box(3)

# ===== SECTION 5 =====
heading2("Section 5: Triggers, Reactions & The Gap")
instruction("Check any that apply. Add as much detail as you like.")

checkbox_line(
    " (disrespect, confrontation, someone difficult)",
    bold_prefix="A triggering situation happened",
)
instruction("What happened? Did you notice the gap between trigger and response?")
fill_line(2)

checkbox_line(
    " (saw someone's anger as their struggle, not an attack on you)",
    bold_prefix="I used the rescuer-helper reframe",
)
instruction("Describe:")
fill_line(1)

checkbox_line(
    ' (e.g., "asshole," "stop being so judgmental")',
    bold_prefix="I caught myself self-labeling",
)
instruction("What did you do when you noticed?")
fill_line(1)

checkbox_line(
    "",
    bold_prefix="Moments of unexpected calm or steadiness",
)
instruction("When? What were you doing?")
fill_line(1)

checkbox_line(
    "",
    bold_prefix='The "armor/shield" feeling from the tribe circle returned',
)
instruction("When? How strong (1\u201310)?")
fill_line(1)

checkbox_line(
    "",
    bold_prefix="I noticed the gap between trigger and response",
)
instruction("What was the trigger? What did you choose to do in the gap?")
fill_line(1)

# ===== SECTION 6 =====
heading2("Section 6: Self-Talk Awareness")
instruction(
    "One of your homework items is to notice how you talk to yourself. "
    "You do not need to change it yet \u2014 just notice."
)
question("What kinds of things do you say to yourself when something goes wrong?")
fill_box(2)
question("Did you notice any shift in your inner voice compared to before our sessions?")
checkbox_line("Yes \u2014 describe: ")
fill_line(1)
checkbox_line("Not really / not sure yet")
checkbox_line("I did not pay attention to this one")

# ===== SECTION 7 =====
heading2("Section 7: Dreams")
instruction(
    "Vivid or unusual dreams after a session are normal \u2014 your subconscious is "
    "processing and clearing. Just jot down keywords."
)
styled_table(
    ["Day", "Dream? (Y/N)", "Brief description (keywords are fine)"],
    [
        ["Day 1", "", ""],
        ["Day 2", "", ""],
        ["Day 3", "", ""],
        ["Day 4", "", ""],
        ["Day 5", "", ""],
        ["Day 6", "", ""],
        ["Day 7", "", ""],
    ],
    col_widths=[2, 2.5, 11],
)
question("Did the fighting / revenge nightmares change at all?")
checkbox_line("Gone or much less frequent")
checkbox_line("Still happening but less intense")
checkbox_line("About the same")
checkbox_line("Worse or new disturbing dreams appeared")
fill_line(1)

# ===== SECTION 8 =====
heading2("Section 8: Session 2 Feedback")
question("What was the most meaningful or memorable part of our session?")
fill_box(2)
question("Was there anything that felt uncomfortable or that you did not understand?")
fill_box(2)
question("Is there anything from Session 2 you want to revisit or go deeper into?")
fill_box(2)

# ===== SECTION 9 =====
heading2("Section 9: Looking Ahead to Session 3")
instruction(
    "In our next session we will work on Circle Therapy \u2014 systematic desensitization "
    "of your specific triggers while in deep relaxation. "
    "Your tribe circle energy and anchor will be part of that protective foundation."
)
question("How ready do you feel for our next session?")
for label in [
    "1 \u2014 Not Ready",
    "2 \u2014 Uneasy",
    "3 \u2014 Neutral",
    "4 \u2014 Ready",
    "5 \u2014 Very Ready",
]:
    checkbox_line(label)
question("If less than 4, what would help you feel more ready?")
fill_line(1)
question("Anything you want to ask, discuss, or bring up next time?")
fill_box(2)

# ===== SAFETY REMINDER =====
heading2("Reminder")
p = doc.add_paragraph()
run = p.add_run(
    "If anything feels overwhelming between sessions:\n"
    "1. Use your anchor: thumb + forefinger, right hand. One breath. Let the steadiness return.\n"
    "2. Email me: jeeth@vrishihypno.com (I respond within 24 hours).\n"
    "3. If urgent or in crisis: call or text 988 (Suicide & Crisis Lifeline, 24/7)."
)
run.font.size = Pt(9.5)
run.font.color.rgb = RGBColor(0x88, 0x66, 0x00)

# ===== FOOTER =====
doc.add_paragraph()
footer = doc.add_paragraph()
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run(
    "VRishi Hypnotherapy  |  T.M. \u2014 Feedback & Self-Tracking After Session 2  |  "
    "CONFIDENTIAL  |  Jithendran Sellamuthu, C.MH., AHA #007913"
)
run.font.size = Pt(7)
run.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUT))
print("Saved:", OUT)
