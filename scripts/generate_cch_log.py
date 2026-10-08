#!/usr/bin/env python3
"""Generate HMI CCH Log as Word document from client JSON + session data."""
import json
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def generate_cch_log(client_json_path, session_idx=0, output_dir=None):
    with open(client_json_path, "r", encoding="utf-8") as f:
        client = json.load(f)

    if not client.get("sessions"):
        print("Error: no sessions in client data")
        sys.exit(1)

    s = client["sessions"][session_idx]
    soap = s.get("soap", {})

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Title
    title = doc.add_heading("HMI Client Contact Hours (CCH) Log", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("")

    # Session info table
    info = doc.add_table(rows=4, cols=4)
    info.style = "Light Grid Accent 1"
    data = [
        ("Student", "Jithendran Sellamuthu", "AHA #", "007913"),
        ("Client Initials", client.get("initials", client["name"]), "Session #", "{} ({})".format(s["num"], "Intake" if s["num"] == 1 else "Follow-up")),
        ("Date", s["date"], "Duration", "{} minutes".format(s.get("duration_min", 60))),
        ("Modality", "Zoom (Virtual)", "Fee", "Pro Bono (CSP)" if client.get("fee") == "pro_bono" else "$250/60min"),
    ]
    for i, row_data in enumerate(data):
        for j, val in enumerate(row_data):
            info.rows[i].cells[j].text = val

    doc.add_paragraph("")

    # 1. Presenting Issue
    doc.add_heading("1. Presenting Issue", level=2)
    doc.add_paragraph(
        "{initials} is a {age}-year-old {occ}presenting with {presenting}. "
        "{history}".format(
            initials=client.get("initials", client["name"]),
            age=client["age"],
            occ="{}, ".format(client["occupation"]) if client.get("occupation") else "",
            presenting=client["presenting"].lower(),
            history=soap.get("subjective", client.get("history", ""))
        )
    )

    # 2. Preparation
    doc.add_heading("2. What preparation did you do for this session?", level=2)
    doc.add_paragraph(
        "Prepared a comprehensive session script following the Kappasinian First Session Flowchart "
        "with client-specific adaptations. Preparation included:"
    )
    prep_items = [
        "Reviewed client intake registration from HMI CSP program",
        "Prepared L.O.V.E. intake structure (Listen, Observe, Verify, Empathize) with appropriate open-ended and closed-ended questions",
        "Prepared screening questions (active care, medications, safety, trauma scope, triggers, panic/dissociation)",
        "Prepared 36-question suggestibility test with both Physical and Emotional EP delivery branches",
        "Prepared Theory of Mind script incorporating client-relevant framing",
        "Prepared pre-induction speech (14 points)",
        "Prepared induction scripts for both EP lanes",
        "Prepared SB 577 Disclosure and Recording Consent aligned with HMI CSP terms",
        "All documents sent to client prior to session",
    ]
    for item in prep_items:
        doc.add_paragraph(item, style="List Bullet")

    # 3. Techniques
    doc.add_heading("3. What hypnotic techniques and suggestions did you use in this session?", level=2)
    doc.add_paragraph("The following techniques were employed:")
    techniques = s.get("techniques", "").split(", ")
    for tech in techniques:
        if tech.strip():
            doc.add_paragraph(tech.strip(), style="List Bullet")

    if soap.get("objective"):
        doc.add_paragraph("")
        doc.add_paragraph("Objective observations: " + soap["objective"])

    # 4. Highlights
    doc.add_heading("4. What were the highlights of this session?", level=2)
    highlights = []
    if client.get("ep"):
        highlights.append(
            "Suggestibility testing revealed {} profile, establishing the delivery lane for all future sessions.".format(client["ep"])
        )
    if soap.get("assessment"):
        highlights.append("Clinical assessment: " + soap["assessment"])
    if s.get("feedback"):
        highlights.append("Client feedback: " + s["feedback"])

    # Add session-specific highlights based on notes
    notes = client.get("notes", {})
    for key, val in notes.items():
        if key not in ("client_preference", "csp_compliance"):
            highlights.append(val)

    for item in highlights:
        doc.add_paragraph(item, style="List Bullet")

    # 5. Supervisor
    doc.add_heading("5. Did you discuss this session with a supervisor and if so what guidance did you receive?", level=2)
    doc.add_paragraph(
        "Prior to this session, reviewed relevant HMI case conference materials. "
        "Plan to discuss this specific session with supervisor to review clinical decisions made during the session."
    )

    # 6. Afterthoughts
    doc.add_heading("6. What afterthoughts and/or conclusions did you have after review of the session?", level=2)
    if soap.get("plan"):
        doc.add_paragraph("Treatment plan going forward: " + soap["plan"])

    doc.add_paragraph("")

    # Footer
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Jithendran Sellamuthu, C.MH. | AHA #007913 | VRishi Hypnotherapy")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(128, 128, 128)
    footer2 = doc.add_paragraph()
    footer2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = footer2.add_run("HMI Community Service Program | Client Initials Only")
    run2.font.size = Pt(9)
    run2.font.color.rgb = RGBColor(128, 128, 128)

    # Save
    if output_dir:
        out = Path(output_dir)
    else:
        out = Path("Clients") / client.get("full_name", client["name"]) / "Documents"
    out.mkdir(parents=True, exist_ok=True)

    filename = "CCH-Log-S{}-{}.docx".format(s["num"], s["date"])
    filepath = out / filename
    doc.save(str(filepath))
    print("CCH Log saved: {}".format(filepath))
    return filepath


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python generate_cch_log.py data/clients/TM-002.json [--session 0] [--output-dir path]")
        sys.exit(1)

    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("client_json")
    parser.add_argument("--session", "-s", type=int, default=0)
    parser.add_argument("--output-dir", "-o")
    args = parser.parse_args()
    generate_cch_log(args.client_json, args.session, args.output_dir)
